"""
Study intake (Phase 6): what it agrees to read, and when it lets go of a source.

Two things here are worth more than the rest. The privacy test, because intake
is the one place where *script code* reads a file straight into a prompt — the
agent never calls Read, so the PreToolUse guard that protects every other path
is structurally blind to this one. And the retention tests, because the first
real run cleared the drop folder after the model finished but before the note
landed, and a source deleted with nothing to show for it is the one failure that
loses work rather than postponing it.
"""
import asyncio
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "runtime"))
sys.path.insert(0, str(REPO / "interface" / "backend"))

import intake                                   # noqa: E402
import privacy                                  # noqa: E402


def _run(repo, *args):
    return subprocess.run(["git", "-C", str(repo), *args],
                          capture_output=True, text=True, encoding="utf-8")


class IntakeBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self.tmp.name)
        self.vault = self.root / "vault"
        (self.vault / "02-Areas" / "Academics" / "CSE-311").mkdir(parents=True)
        (self.vault / "02-Areas" / "Academics" / "MATH-208").mkdir(parents=True)
        self.drop = self.vault / "00-Inbox" / "intake"
        (self.drop / "CSE-311").mkdir(parents=True)

        (self.vault / ".gitignore").write_text("secret/\n", encoding="utf-8")
        _run(self.vault, "init", "-b", "master")
        _run(self.vault, "config", "user.name", "T")
        _run(self.vault, "config", "user.email", "t@e.com")
        _run(self.vault, "add", "-A")
        _run(self.vault, "commit", "-m", "seed")

        self._saved = (intake.VAULT, intake.DROP, intake.ATTACH,
                       intake.COURSES_ROOT, privacy._RUNTIME)
        intake.VAULT = self.vault
        intake.DROP = self.drop
        intake.ATTACH = self.vault / "99-Meta" / "Attachments"
        intake.COURSES_ROOT = self.vault / "02-Areas" / "Academics"
        privacy._RUNTIME = self.root
        (self.root / "privacy.config.json").write_text(
            json.dumps({"model_allow": []}), encoding="utf-8")
        self._bust()

    def _bust(self):
        privacy.model_allow_raw.cache_clear()
        privacy.model_allow_prefixes.cache_clear()
        privacy.VaultPrivacy._git_ignored.cache_clear()

    def tearDown(self):
        (intake.VAULT, intake.DROP, intake.ATTACH,
         intake.COURSES_ROOT, privacy._RUNTIME) = self._saved
        self._bust()
        self.tmp.cleanup()

    def drop_file(self, rel: str, body: str = "x" * 500) -> Path:
        p = self.drop / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(body, encoding="utf-8")
        return p


class TestScan(IntakeBase):
    def test_a_course_subfolder_is_picked_up(self):
        self.drop_file("CSE-311/week-4.md")
        items, problems = intake.scan()
        self.assertEqual([(p.name, c) for p, c in items], [("week-4.md", "CSE-311")])
        self.assertEqual(problems, [])

    def test_course_folder_matching_is_case_insensitive_but_files_use_the_real_name(self):
        self.drop_file("cse-311/week-4.md")
        items, _ = intake.scan()
        self.assertEqual(items[0][1], "CSE-311")   # not "cse-311"

    def test_an_unknown_course_is_reported_never_guessed(self):
        """The contract's filing rule: do not invent a location."""
        self.drop_file("CSE-999/mystery.md")
        items, problems = intake.scan()
        self.assertEqual(items, [])
        self.assertTrue(any("no such course folder" in why for _, why in problems))

    def test_a_top_level_file_is_reported(self):
        self.drop_file("loose.md")
        items, problems = intake.scan()
        self.assertEqual(items, [])
        self.assertTrue(any("top level" in why for _, why in problems))

    def test_an_unsupported_type_is_reported(self):
        # .docx, not .pptx — slides became supported on 2026-07-31 and this
        # fixture was the first thing to notice.
        self.drop_file("CSE-311/essay.docx")
        items, problems = intake.scan()
        self.assertEqual(items, [])
        self.assertTrue(any("unsupported type" in why for _, why in problems))


class TestPrivacyBoundary(IntakeBase):
    """Intake reads files into a prompt itself, so it must apply the model
    boundary itself. Nothing else in the system would catch this."""

    def test_a_gitignored_source_is_refused(self):
        (self.vault / ".gitignore").write_text(
            "secret/\n00-Inbox/intake/CSE-311/\n", encoding="utf-8")
        self._bust()
        self.drop_file("CSE-311/private.md")
        items, problems = intake.scan()
        self.assertEqual(items, [], "a gitignored file reached the intake queue")
        self.assertTrue(any("refused at the model boundary" in why
                            for _, why in problems), problems)

    def test_an_exempted_gitignored_source_is_allowed(self):
        """Option B's exemption applies here too — if the model may read it,
        intake may feed it."""
        (self.vault / ".gitignore").write_text(
            "00-Inbox/intake/CSE-311/\n", encoding="utf-8")
        (self.root / "privacy.config.json").write_text(
            json.dumps({"model_allow": ["00-Inbox/intake/CSE-311/"]}), encoding="utf-8")
        self._bust()
        self.drop_file("CSE-311/allowed.md")
        items, problems = intake.scan()
        self.assertEqual(len(items), 1, problems)


class TestExtract(IntakeBase):
    def test_long_text_is_truncated_and_says_so(self):
        p = self.drop_file("CSE-311/big.txt", "y" * (intake.TEXT_BUDGET + 5000))
        text, truncated = intake.extract(p)
        self.assertTrue(truncated)
        self.assertEqual(len(text), intake.TEXT_BUDGET)

    def test_short_text_is_not_flagged(self):
        p = self.drop_file("CSE-311/small.txt", "z" * 300)
        _, truncated = intake.extract(p)
        self.assertFalse(truncated)


def _has_pptx():
    try:
        import pptx  # noqa: F401
        return True
    except ImportError:
        return False


@unittest.skipUnless(_has_pptx(), "python-pptx not installed")
class TestPptx(IntakeBase):
    """Lecture slides. 49 of AA-210's 125 source documents are decks, so this
    path carries a large share of the course's material."""

    def deck(self, name="d.pptx", slides=3, footer="Copyright (c) 2020"):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        for i in range(slides):
            s = prs.slides.add_slide(prs.slide_layouts[5])   # title only
            s.shapes.title.text = f"Topic {i + 1}"
            box = s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(3))
            tf = box.text_frame
            tf.text = f"point {i + 1}a"
            tf.add_paragraph().text = f"point {i + 1}b"
            if footer:                       # master-slide furniture on every slide
                tf.add_paragraph().text = footer
            s.notes_slide.notes_text_frame.text = f"note {i + 1}"
        p = self.drop / "CSE-311" / name
        p.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(p))
        return p

    def test_slide_structure_is_kept(self):
        text = intake.extract_pptx(self.deck())
        self.assertIn("### Slide 1 — Topic 1", text)
        self.assertIn("### Slide 3 — Topic 3", text)
        self.assertIn("- point 2a", text)

    def test_speaker_notes_are_included(self):
        """Often the only place a deck explains itself rather than listing."""
        text = intake.extract_pptx(self.deck())
        self.assertIn("> Speaker notes: note 2", text)

    def test_repeated_footer_is_stripped(self):
        text = intake.extract_pptx(self.deck(slides=6))
        self.assertNotIn("Copyright", text)
        self.assertIn("- point 1a", text)     # real content survives

    def test_a_line_on_only_one_slide_is_content_not_furniture(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        for i in range(6):
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = f"T{i}"
            tf = s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)).text_frame
            tf.text = "shared footer"
            if i == 0:
                tf.add_paragraph().text = "a unique and important claim"
        p = self.drop / "CSE-311" / "u.pptx"
        prs.save(str(p))
        text = intake.extract_pptx(p)
        self.assertNotIn("shared footer", text)
        self.assertIn("a unique and important claim", text)

    def test_pptx_is_accepted_by_the_scanner(self):
        self.deck(name="lecture.pptx")
        items, problems = intake.scan()
        self.assertEqual([p.name for p, _ in items], ["lecture.pptx"], problems)

    def test_extract_routes_pptx_and_reports_length(self):
        text, truncated = intake.extract(self.deck())
        self.assertFalse(truncated)
        self.assertIn("### Slide 1", text)


class TestSlideEquations(unittest.TestCase):
    """Equations are OMML, and python-pptx's `.text` walks `a:t` runs only.

    AA-210's decks carry 4,710 of them; before this, three slides of the
    friction deck extracted completely empty because the whole slide *is* the
    formula. These are pure XML functions, so they need no deck on disk.
    """

    def frag(self, xml: str):
        from lxml import etree
        ns = ('xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
              'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
              'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
              'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"')
        return etree.fromstring(xml.replace("<root>", f"<root {ns}>"))

    def math(self, body: str) -> str:
        return intake._omml(self.frag(f"<root><m:oMath>{body}</m:oMath></root>")[0])

    def test_subscript_reads_as_underscore(self):
        got = self.math("<m:r><m:t>F=</m:t></m:r>"
                        "<m:sSub><m:e><m:r><m:t>μ</m:t></m:r></m:e>"
                        "<m:sub><m:r><m:t>s</m:t></m:r></m:sub></m:sSub>"
                        "<m:r><m:t>N</m:t></m:r>")
        self.assertEqual(got, "F=μ_sN")

    def test_fraction_is_parenthesised_on_both_sides(self):
        """`Ph/W` would be wrong; the slide means the whole numerator."""
        got = self.math("<m:f><m:num><m:r><m:t>Ph</m:t></m:r></m:num>"
                        "<m:den><m:r><m:t>W</m:t></m:r></m:den></m:f>")
        self.assertEqual(got, "(Ph)/(W)")

    def test_radical_and_superscript(self):
        self.assertEqual(
            self.math("<m:rad><m:e><m:r><m:t>x</m:t></m:r></m:e></m:rad>"),
            "sqrt(x)")
        self.assertEqual(
            self.math("<m:sSup><m:e><m:r><m:t>r</m:t></m:r></m:e>"
                      "<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>"),
            "r^2")

    def test_run_properties_inside_math_are_not_text(self):
        """`m:r` carries an `a:rPr` sibling of `m:t` — formatting, not content."""
        got = self.math('<m:r><a:rPr lang="en-US" i="1"/><m:t>M</m:t></m:r>')
        self.assertEqual(got, "M")

    def test_unknown_construct_degrades_to_its_children(self):
        got = self.math("<m:borderBox><m:e><m:r><m:t>Q</m:t></m:r></m:e></m:borderBox>")
        self.assertEqual(got, "Q")

    def test_math_keeps_its_place_in_the_sentence(self):
        """'it is NOT true that F=μN' reverses if the math is appended after."""
        p = self.frag(
            "<root><a:p>"
            "<a:r><a:t>it is NOT true that </a:t></a:r>"
            "<m:oMath><m:r><m:t>F=μN</m:t></m:r></m:oMath>"
            "</a:p></root>")[0]
        self.assertEqual(intake._para_text(p), "it is NOT true that F=μN")

    def test_math_italic_unicode_is_normalised(self):
        """Decks use Cambria Math's italic plane; U+1D439 is an F, not a glyph."""
        p = self.frag("<root><a:p><m:oMath><m:r><m:t>\U0001d439</m:t></m:r>"
                      "</m:oMath></a:p></root>")[0]
        self.assertEqual(intake._para_text(p), "F")

    def test_walk_descends_into_alternate_content(self):
        """The wrapper PowerPoint puts around every equation-bearing shape."""
        tree = self.frag(
            "<root><p:spTree>"
            "<mc:AlternateContent>"
            "  <mc:Choice><p:sp><p:txBody><a:p><a:r><a:t>live</a:t></a:r>"
            "  </a:p></p:txBody></p:sp></mc:Choice>"
            "  <mc:Fallback><p:sp><p:txBody><a:p><a:r><a:t>flattened</a:t>"
            "  </a:r></a:p></p:txBody></p:sp></mc:Fallback>"
            "</mc:AlternateContent>"
            "</p:spTree></root>")[0]
        found = [intake._body_text(sp.find(intake._PML_NS + "txBody"))
                 for sp in intake._walk_shapes(tree)]
        self.assertEqual(found, ["live"])

    def test_walk_flattens_groups(self):
        tree = self.frag(
            "<root><p:spTree><p:grpSp>"
            "<p:sp><p:txBody><a:p><a:r><a:t>grouped</a:t></a:r></a:p>"
            "</p:txBody></p:sp>"
            "</p:grpSp></p:spTree></root>")[0]
        self.assertEqual(len(list(intake._walk_shapes(tree))), 1)


class TestSourceRetention(IntakeBase):
    """When may intake clear the drop folder? Only once a note actually reached
    the vault. Proposing is not landing."""

    def _fake_run(self, result, applied):
        async def fake_intake_one(path, course, model="sonnet"):
            return result
        import applier
        self._orig = (intake.intake_one, applier.apply_run)
        intake.intake_one = fake_intake_one
        applier.apply_run = lambda files, actor: applied
        self.addCleanup(self._restore)

    def _restore(self):
        import applier
        intake.intake_one, applier.apply_run = self._orig

    def test_source_is_cleared_when_a_note_lands(self):
        src = self.drop_file("CSE-311/week-4.md")
        self._fake_run({"ok": True, "proposals": 1, "files": ["p.md"], "summary": ""},
                       [{"action": "create", "target": "x.md", "reason": None}])
        intake.run()
        self.assertFalse(src.exists())

    def test_source_survives_when_applying_fails(self):
        """The regression: the first real run proposed a note, could not apply
        it because git hung, and cleared the source anyway."""
        src = self.drop_file("CSE-311/week-4.md")
        self._fake_run({"ok": True, "proposals": 1, "files": ["p.md"], "summary": ""},
                       [])
        rc = intake.run()
        self.assertTrue(src.exists(), "source was cleared with nothing applied")
        self.assertEqual(rc, 1)

    def test_source_survives_when_the_model_proposes_nothing(self):
        src = self.drop_file("CSE-311/week-4.md")
        self._fake_run({"ok": True, "proposals": 0, "files": [], "summary": ""}, [])
        intake.run()
        self.assertTrue(src.exists())

    def test_declining_thin_material_is_not_a_failure(self):
        """The brief tells the model that material too thin to work with should
        produce nothing. Exiting non-zero for that made the code contradict its
        own instructions — the palette renders it as '✗ failed'."""
        self.drop_file("CSE-311/week-4.md")
        self._fake_run({"ok": True, "proposals": 0, "files": [], "summary": ""}, [])
        self.assertEqual(intake.run(), 0)

    def test_a_real_failure_still_exits_non_zero(self):
        self.drop_file("CSE-311/week-4.md")
        self._fake_run({"ok": False, "proposals": 0, "files": [],
                        "error": "timed out"}, [])
        self.assertEqual(intake.run(), 1)

    def test_a_held_proposal_still_counts_as_landed(self):
        """A hold means the change is staged and recorded, not lost."""
        src = self.drop_file("CSE-311/week-4.md")
        self._fake_run({"ok": True, "proposals": 1, "files": ["p.md"], "summary": ""},
                       [{"action": "held", "target": "x.md", "reason": "already exists"}])
        intake.run()
        self.assertFalse(src.exists())

    def test_keep_leaves_the_source_alone(self):
        src = self.drop_file("CSE-311/week-4.md")
        self._fake_run({"ok": True, "proposals": 1, "files": ["p.md"], "summary": ""},
                       [{"action": "create", "target": "x.md", "reason": None}])
        intake.run(keep=True)
        self.assertTrue(src.exists())


# --------------------------------------------------------------------------
# where the document rides
# --------------------------------------------------------------------------

# Windows' CreateProcess ceiling. The SDK passes the system prompt — where the
# brief goes — as `--append-system-prompt` on the command line, so a brief
# carrying the document blows straight past this.
CMDLINE_MAX = 32767


class TestTheDocumentRidesStdin(IntakeBase):
    """The document goes in `material`, never in `brief`.

    Every intake of a real lecture PDF failed on this: a 60,000-character
    extraction embedded in the brief made the spawn fail with
    `CLINotFoundError: Claude Code not found`, naming a bundled binary that was
    present and ran fine by hand. The error names the wrong thing entirely,
    which is why it read as a broken install rather than an oversized argument.
    `fleet.run_one` already had `material` for exactly this and its docstring
    documents this precise failure; intake was the caller that never used it.
    """

    def test_the_brief_never_carries_the_document(self):
        body = "SENTINEL-BODY " * 5000
        brief = intake.build_brief(Path("topic01-logic.pdf"), "CSE-311", False)
        self.assertNotIn("SENTINEL-BODY", brief)
        self.assertIn("SENTINEL-BODY", intake.build_material(body))

    def test_the_brief_does_not_grow_with_the_document(self):
        """60k of extracted slides is the ordinary case here, not the extreme."""
        brief = intake.build_brief(Path("topic01-logic.pdf"), "CSE-311", False)
        orientation = f"{intake.RULES}\n\n## Your brief\n\n{brief}"
        self.assertLess(len(orientation), CMDLINE_MAX)
        self.assertLess(len(intake.build_brief(Path("x.pdf"), "CSE-311", True)),
                        CMDLINE_MAX // 4)

    def test_intake_one_hands_the_text_to_run_one_as_material(self):
        """The wiring, not just the shapes: a build_material nothing passes to
        run_one would leave the document unread and the notes empty."""
        src = self.drop_file("CSE-311/week-4.md")
        src.write_text("# Week 4\n\n" + ("real lecture prose. " * 60),
                       encoding="utf-8")
        seen = {}

        async def fake_run_one(spec, timeout_s=420, model_override=None,
                               rules=None, material=None):
            seen["brief"], seen["material"] = spec.brief, material
            return {"ok": True, "proposals": 0, "files": [], "summary": ""}

        import fleet as fl
        orig, fl.run_one = fl.run_one, fake_run_one
        try:
            asyncio.run(intake.intake_one(src, "CSE-311"))
        finally:
            fl.run_one = orig

        self.assertIsNotNone(seen["material"], "the document was not passed at all")
        self.assertIn("real lecture prose", seen["material"])
        self.assertNotIn("real lecture prose", seen["brief"])


class TestContinuation(IntakeBase):
    """Finishing a source that ran past one pass's budget.

    The budget is a ceiling on a *pass*, not on a document — otherwise a long
    deck is silently abandoned at 75% and the notes look complete. Seven CSE-311
    decks came in at 89% coverage before this existed.
    """

    def setUp(self):
        super().setUp()
        self.cov = self.root / "intake.state.json"
        self.attach = self.vault / "99-Meta" / "Attachments"
        self.attach.mkdir(parents=True, exist_ok=True)
        self._budget = intake.TEXT_BUDGET
        intake.TEXT_BUDGET = 100
        self.addCleanup(setattr, intake, "TEXT_BUDGET", self._budget)

    def source(self, name, chars, course="CSE-311", cited=True):
        (self.attach / name).write_text("A" * chars, encoding="utf-8")
        if cited:
            # One note per source: a shared filename made the second call
            # overwrite the first, so only the last source looked cited.
            note = (self.vault / "02-Areas" / "Academics" / course / "lectures"
                    / f"{Path(name).stem}-note.md")
            note.parent.mkdir(parents=True, exist_ok=True)
            note.write_text(f"# N\n\n> Source: ![[{name}]]\n", encoding="utf-8")
        return self.attach / name

    def test_a_pass_reads_the_next_budget_not_the_first_one_again(self):
        full = "".join(str(i % 10) for i in range(250))
        first, t1 = intake.slice_of(full, 0)
        second, t2 = intake.slice_of(full, 100)
        third, t3 = intake.slice_of(full, 200)
        self.assertEqual((len(first), t1), (100, True))
        self.assertEqual((len(second), t2), (100, True))
        self.assertEqual((len(third), t3), (50, False))
        self.assertEqual(first + second + third, full)

    def test_only_sources_with_unread_material_are_pending(self):
        self.source("long.txt", 250)
        self.source("short.txt", 60)
        keys = [p["key"] for p in intake.pending("CSE-311", self.cov)]
        self.assertIn("CSE-311/long.txt", keys)
        self.assertNotIn("CSE-311/short.txt", keys)

    def test_a_source_with_no_notes_is_not_this_courses_problem(self):
        """Attachments are a flat folder shared by every course; a deck belongs
        to the course whose notes cite it."""
        self.source("orphan.txt", 250, cited=False)
        self.assertEqual(intake.pending("CSE-311", self.cov), [])

    def test_coverage_with_no_record_is_seeded_from_history(self):
        """The old code always read exactly TEXT_BUDGET, so that is the honest
        default — and it means the first --continue needs no migration."""
        self.source("long.txt", 250)
        p = intake.pending("CSE-311", self.cov)[0]
        self.assertEqual((p["read_to"], p["total"]), (100, 250))

    def test_a_recorded_pass_resumes_from_where_it_stopped(self):
        self.source("long.txt", 250)
        self.cov.write_text(json.dumps(
            {"CSE-311/long.txt": {"read_to": 200, "total": 250}}), encoding="utf-8")
        p = intake.pending("CSE-311", self.cov)[0]
        self.assertEqual(p["read_to"], 200)

    def test_a_fully_read_source_disappears_from_pending(self):
        self.source("long.txt", 250)
        self.cov.write_text(json.dumps(
            {"CSE-311/long.txt": {"read_to": 250, "total": 250}}), encoding="utf-8")
        self.assertEqual(intake.pending("CSE-311", self.cov), [])

    def test_a_failed_pass_does_not_advance_coverage(self):
        """Otherwise the part nobody read is skipped forever — the same rule
        that keeps a source in the drop folder when its notes did not land."""
        self.source("long.txt", 250)

        async def fake(*a, **k):
            return {"ok": False, "proposals": 0, "files": [], "error": "boom"}

        orig, intake.continue_one = intake.continue_one, fake
        try:
            intake.run_continue("CSE-311", cov_path=self.cov)
        finally:
            intake.continue_one = orig
        self.assertFalse(self.cov.exists() and json.loads(
            self.cov.read_text(encoding="utf-8")))

    def test_a_pass_that_found_nothing_new_still_advances(self):
        """A section that genuinely adds nothing is a correct outcome; not
        advancing would loop on it forever."""
        self.source("long.txt", 250)

        async def fake(src, course, start, total, existing, model="sonnet"):
            return {"ok": True, "proposals": 0, "files": [], "summary": "",
                    "read_to": start + 100, "total": total}

        orig, intake.continue_one = intake.continue_one, fake
        try:
            intake.run_continue("CSE-311", cov_path=self.cov)
        finally:
            intake.continue_one = orig
        rec = json.loads(self.cov.read_text(encoding="utf-8"))["CSE-311/long.txt"]
        self.assertEqual(rec["read_to"], 200)

    def test_the_continuation_brief_names_what_already_exists(self):
        brief = intake.build_continue_brief(
            Path("topic02-proofs.pdf"), "CSE-311", 60000, 79793, 79793,
            ["02-Areas/Academics/CSE-311/lectures/mathematical-induction.md"])
        self.assertIn("continuation", brief.lower())
        self.assertIn("60,000", brief)
        self.assertIn("mathematical-induction.md", brief)
        # still an instruction sheet, not a document — the ceiling still applies
        self.assertLess(len(f"{intake.RULES}\n\n{brief}"), CMDLINE_MAX)


if __name__ == "__main__":
    unittest.main()
